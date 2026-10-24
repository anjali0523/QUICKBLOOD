import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml

def create_quickblood_document(output_path):
    doc = docx.Document()

    # Set page setup - Letter (8.5 x 11 in) with 0.38" top/bottom, 0.5" left/right
    for section in doc.sections:
        section.top_margin = Inches(0.38)
        section.bottom_margin = Inches(0.38)
        section.left_margin = Inches(0.50)
        section.right_margin = Inches(0.50)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Color definitions
    COLOR_TITLE = RGBColor(163, 22, 28)       # Deep Crimson #A3161C
    COLOR_HEADING = RGBColor(140, 18, 24)     # Dark Red #8C1218
    COLOR_BODY = RGBColor(30, 33, 40)         # Charcoal Dark Text #1E2128
    COLOR_KEYWORD = RGBColor(45, 50, 60)

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=1.5, line_spacing=1.03):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.0)
            r.font.color.rgb = COLOR_BODY
        return p

    def add_heading(text, space_before=2.5, space_after=1.0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.02
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = COLOR_HEADING
        return p

    # --- TITLE ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(1.5)
    p_title.paragraph_format.line_spacing = 1.02
    p_title.paragraph_format.keep_with_next = True

    r_title = p_title.add_run("QUICKBLOOD: SMART REAL-TIME BLOOD DONOR & EMERGENCY RESPONSE SYSTEM")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(12.0)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_TITLE

    # Thin Crimson Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(2.0)
    p_div.paragraph_format.line_spacing = 1.0
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="A3161C"/>'
                     r'</w:pBdr>')
    p_div._p.get_or_add_pPr().append(pBdr)

    # --- ABSTRACT ---
    add_heading("Abstract", space_before=1.5, space_after=1.0)

    p_abs1 = add_p(space_after=1.5)
    p_abs1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_abs1.add_run(
        "Access to compatible blood becomes critical when individuals encounter acute trauma, emergency surgeries, or sudden clinical shortages "
        "where immediate and reliable assistance is required. Existing blood donation solutions often address isolated aspects such as "
        "periodic donor drives, physical blood bank registries, or unverified social media appeals, resulting in fragmented and delayed life-saving support. "
        "QUICKBLOOD is a smart web-based blood donor matching and emergency response system designed to integrate donor discovery, compatibility "
        "intelligence, and emergency assistance into a unified platform."
    )

    p_abs2 = add_p(space_after=1.5)
    p_abs2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_abs2.add_run(
        "The system provides an Instant Donor Search Engine that matches recipient requisitions with nearby registered donors using location-based "
        "information, ABO/Rh blood groups, and live availability status. It also supports Blood Compatibility Guidance, enabling users to remain "
        "aware of biological transferability and eligibility criteria. In emergency situations, healthcare personnel and caregivers can broadcast "
        "urgent hospital requisitions classified by clinical urgency, which helps reduce procurement delays and facilitates rapid access to verified voluntary donors."
    )

    p_abs3 = add_p(space_after=2.0)
    p_abs3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_abs3.add_run(
        "QUICKBLOOD also incorporates a Request Lifecycle Ledger for recording and tracking emergency blood requests through fulfillment, and supports "
        "real-time donor availability management to maintain accurate donor readiness. The proposed system is implemented as a full-stack web application "
        "using Python Flask, SQLAlchemy ORM, SQLite relational storage, and responsive interface mechanisms. By combining proactive voluntary donor "
        "coordination with rapid emergency response, QUICKBLOOD aims to provide a more reliable and comprehensive approach to emergency blood procurement."
    )

    # --- KEYWORDS ---
    add_heading("Keywords", space_before=1.5, space_after=1.0)
    p_kw = add_p(space_after=2.0)
    p_kw.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_kw = p_kw.add_run(
        "Blood Donation, Emergency Response, Donor Matching, Blood Compatibility, Urgency Triage, Healthcare Informatics, Flask Application, Real-Time Availability, Request Ledger, Geolocation Search"
    )
    r_kw.font.italic = True
    r_kw.font.color.rgb = COLOR_KEYWORD

    # --- METHODOLOGY ---
    add_heading("Methodology", space_before=1.5, space_after=1.0)

    p_meth_intro = add_p(space_after=1.5)
    p_meth_intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_meth_intro.add_run(
        "The methodology of QUICKBLOOD follows a layered approach that integrates donor intelligence, compatibility logic, urgency triage, "
        "emergency response, and relational data storage into a single Flask-based system."
    )

    def add_step(title, desc_parts):
        """Adds a numbered step with custom styled description parts."""
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_t.paragraph_format.space_before = Pt(1.5)
        p_t.paragraph_format.space_after = Pt(0.5)
        p_t.paragraph_format.line_spacing = 1.02
        p_t.paragraph_format.keep_with_next = True
        
        r_t = p_t.add_run(title)
        r_t.font.name = 'Calibri'
        r_t.font.size = Pt(8.2)
        r_t.bold = True
        r_t.font.color.rgb = COLOR_HEADING

        p_d = doc.add_paragraph()
        p_d.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(1.5)
        p_d.paragraph_format.line_spacing = 1.02

        for part in desc_parts:
            text = part.get('text', '')
            bold = part.get('bold', False)
            color = part.get('color', COLOR_BODY)
            r = p_d.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(7.8)
            r.bold = bold
            r.font.color.rgb = color

    # 1. User Interaction Layer
    add_step(
        "1. User Interaction Layer",
        [{'text': "The process begins with the responsive web application through which users access the major functions of QUICKBLOOD. The application provides interfaces for the Instant Donor Search, Emergency Requests Feed, Blood Compatibility Guide, and User Dashboard."}]
    )

    # 2. User Registration and Role Management
    add_step(
        "2. User Registration and Role Management",
        [{'text': "The system collects user details and configures role-based accounts (Donor, Receiver, or Dual Role) with secure Werkzeug password hashing. Donor profiles capture blood group, municipal location, contact details, and historical donation metrics."}]
    )

    # 3. Blood Compatibility and Matching Logic
    add_step(
        "3. Blood Compatibility and Matching Logic",
        [
            {'text': "Patient blood requirements are processed using transfusion compatibility logic. Donor-recipient compatibility is structured into three clinical categories:\n"},
            {'text': "Universal Donor (O-)", 'bold': True},
            {'text': " – Eligible to provide red blood cells safely to all blood groups in emergency trauma.\n"},
            {'text': "Universal Recipient (AB+)", 'bold': True},
            {'text': " – Capable of receiving red blood cells safely from all ABO and Rh blood groups.\n"},
            {'text': "Specific Compatible Groups", 'bold': True},
            {'text': " – Compatible pairings (A+, A-, B+, B-, O+, AB-) matched to prevent hemolytic transfusion reactions.\n"},
            {'text': "These classifications are displayed to guide users and medical staff toward safe, clinically valid donor selections."}
        ]
    )

    # 4. Location-Based Donor Discovery
    add_step(
        "4. Location-Based Donor Discovery",
        [{'text': "The system applies location-based filtering to identify nearby voluntary donors within the same municipality or metropolitan area. Geographic matching minimizes transit latency between prospective donors and healthcare facilities."}]
    )

    # 5. Emergency Request Broadcast and Dispatch
    add_step(
        "5. Emergency Request Broadcast and Dispatch",
        [{'text': "When an emergency blood shortage occurs, authenticated users initiate an emergency requisition capturing patient name, blood group, hospital facility, units needed, urgency status, and contact phone number."}]
    )

    # 6. Urgency Triage and Prioritization
    add_step(
        "6. Urgency Triage and Prioritization",
        [{'text': "Submitted requests are processed and tagged into clinical urgency tiers (Normal, Urgent, and Critical Emergency). Critical requisitions are surfaced immediately on the public emergency feed with high-visibility alerts to accelerate life-saving response."}]
    )

    # 7. Direct Telephony and Contact Access
    add_step(
        "7. Direct Telephony and Contact Access",
        [{'text': "The system provides direct telephony action triggers (tel: protocol) on matching donor and request listings so assistance can be initiated during an emergency without navigating through intermediary communication channels."}]
    )

    # 8. Real-Time Donor Availability Management
    add_step(
        "8. Real-Time Donor Availability Management",
        [{'text': "Registered donors can dynamically toggle their donation readiness between Available and Unavailable via asynchronous AJAX endpoints, ensuring search results only reflect donors currently eligible and prepared to donate."}]
    )

    # 9. Request Lifecycle and Fulfillment Ledger
    add_step(
        "9. Request Lifecycle and Fulfillment Ledger",
        [{'text': "Each emergency event and its status are recorded in the Request Lifecycle Ledger. Requisition creators can mark requests as Fulfilled upon successful donation, providing structured tracking and maintaining auditability."}]
    )

    # 10. System Integration
    add_step(
        "10. System Integration",
        [
            {'text': "Finally, all components are integrated through a unified architecture consisting of four major layers:\n"},
            {'text': "User Layer → Matching Intelligence Layer → Emergency & Data Layer → Database & Infrastructure Layer\n", 'bold': True, 'color': COLOR_HEADING},
            {'text': "The User Layer handles template rendering and user interaction. The Matching Intelligence Layer performs compatibility validation and location-based filtering. The Emergency & Data Layer manages urgency triage, direct telephony triggers, and live availability toggling. The Database & Infrastructure Layer manages SQLite relational storage, SQLAlchemy ORM mappings, and RESTful API endpoints.\n"
                     "Thus, the overall methodology enables QUICKBLOOD to move from proactive donor coordination to rapid emergency blood fulfillment through one unified platform."}
        ]
    )

    doc.save(output_path)
    print(f"Document saved to: {output_path}")

if __name__ == '__main__':
    out_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'QUICKBLOOD_Project_Document.docx')
    create_quickblood_document(out_file)
